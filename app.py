import sqlite3
from datetime import datetime
from flask import Flask, g, request, render_template_string, redirect, url_for, send_file
import io

try:
    import openpyxl
except ImportError:
    openpyxl = None

app = Flask(__name__)
DB = "responses.db"

# ---------------------------------------------------------------------------
# Questionnaire definition
# ---------------------------------------------------------------------------
TITLE = "Facilities Management Practices and Performance of Office Buildings"
SUBTITLE = "A Study of Office Buildings in Abuja, Federal Capital Territory"

INTRO_TEXT = (
    "This questionnaire is designed purely for academic research purposes, to examine the "
    "relationship between facilities management (FM) practices and the performance of office "
    "buildings in Nigerian organisations, with particular reference to Abuja. You have been "
    "selected to participate based on your role and experience in facilities-related activities "
    "in your office building. Kindly respond as candidly and accurately as possible; there is no "
    "right or wrong answer. All information supplied will be treated with strict confidentiality "
    "and used solely for the purpose of this academic study."
)

# Section A: demographic dropdown questions.
# Each field name becomes a DB column and a form field name.
DEMOGRAPHICS = [
    {
        "field": "gender",
        "label": "Gender",
        "options": ["Male", "Female"],
    },
    {
        "field": "age",
        "label": "Age",
        "options": [
            "Below 25 years",
            "25 - 34 years",
            "35 - 44 years",
            "45 - 54 years",
            "55 years and above",
        ],
    },
    {
        "field": "education",
        "label": "Highest Educational Qualification",
        "options": [
            "OND/HND",
            "Bachelor's Degree",
            "Postgraduate Degree (Master's/PhD)",
            "Professional Certification (e.g., IFMA, RICS, NIQS)",
        ],
    },
    {
        "field": "experience_years",
        "label": "Years of Experience in Facilities Management / Occupancy of Current Building",
        "options": ["Less than 2 years", "2 - 5 years", "6 - 10 years", "Above 10 years"],
    },
    {
        "field": "respondent_category",
        "label": "Respondent Category",
        "options": ["Building User/Tenant", "Facilities Management Staff/Technical Expert"],
    },
    {
        "field": "building_type",
        "label": "Building/Organization Type",
        "options": ["Government Building", "Multi-tenanted Building", "Private Corporate Building"],
    },
]

# Sections B-F: Likert-scale statement blocks.
# Each section has a code (used as a field prefix) and a flat list of items.
# An item can optionally be preceded by a "group" sub-heading (used in Section C).
SECTIONS = [
    {
        "code": "B",
        "title": "SECTION B: PREVAILING FACILITIES MANAGEMENT PRACTICES ADOPTED BY ORGANIZATIONS",
        "items": [
            {"text": "My organization has an established/documented facilities management (FM) policy or framework guiding its facility operations."},
            {"text": "My organization primarily adopts a preventive maintenance approach rather than fixing problems only after they occur (reactive/breakdown-driven maintenance)."},
            {"text": "My organization utilizes computerized or technology-enabled tools (e.g., CMMS, CAFM, BIM) in managing facility operations."},
            {"text": "My organization outsources some or all of its facilities management services to external/professional service providers."},
            {"text": "My organization uses performance measurement tools or metrics to track the delivery of facilities management services."},
            {"text": "My organization provides regular training and capacity development for facilities management personnel."},
            {"text": "My organization prioritizes the provision of a safe environment, satisfactory physical working conditions, and quality service materials in its facilities management practice."},
        ],
    },
    {
        "code": "C",
        "title": "SECTION C: FACILITY MANAGEMENT PRACTICES IN THE OFFICE BUILDING",
        "items": [
            {"group": "Regular Inspections", "text": "Building components and systems in my organization's office building are systematically inspected on a scheduled basis."},
            {"text": "Routine checks are conducted to monitor the general condition of equipment and building systems."},
            {"text": "Detailed condition audits are carried out periodically to comprehensively assess the state of building assets."},
            {"text": "Compliance inspections are conducted to verify adherence to safety and regulatory standards."},
            {"text": "Findings from inspections are documented and systematically recorded (e.g., using checklists, registers, or digital tools)."},
            {"text": "Technology-based tools (e.g., CAFM, CMMS, BIM) are used to schedule, record, or track inspections."},
            {"text": "Inspection data collected is used to inform maintenance prioritization and capital planning decisions."},

            {"group": "Preventive Maintenance", "text": "My organization has clear policies and procedures guiding the planning and scheduling of maintenance activities."},
            {"text": "Maintenance activities are primarily preventive (carried out on a planned schedule) rather than reactive (carried out only after a failure occurs)."},
            {"text": "Predictive maintenance techniques (e.g., condition monitoring) are used to anticipate and pre-empt equipment failure before it occurs."},
            {"text": "Maintenance schedules for building systems (e.g., HVAC, electrical, plumbing) are consistently followed."},
            {"text": "Preventive maintenance practices in my organization have reduced the frequency of building system breakdowns."},

            {"group": "Space Optimization", "text": "Office space allocation is based on functional need rather than organizational hierarchy or seniority."},
            {"text": "My organization monitors space utilization to identify underused or overcrowded areas."},
            {"text": "Office layouts are adapted or reconfigured to accommodate changing work patterns (e.g., flexible or hybrid working)."},
            {"text": "Storage and other spaces are effectively managed to prevent the accumulation of unused furniture or equipment."},

            {"group": "Safety and Compliance", "text": "My organization treats health and safety as an embedded organizational culture, not merely a regulatory compliance requirement."},
            {"text": "Safety infrastructure (e.g., fire extinguishers, safety signage, emergency exits) is regularly maintained and inspected to ensure it remains functional."},
            {"text": "Risk in my organization's building is proactively identified and managed rather than addressed only after an incident occurs."},
            {"text": "My organization maintains vigilance about safety, even in the absence of recent incidents."},

            {"group": "Sustainability Initiatives", "text": "My organization implements energy efficiency programmes within the building (e.g., efficient lighting and equipment)."},
            {"text": "My organization has adopted or explored renewable energy sources to supplement its power supply."},
            {"text": "Waste minimization and water conservation practices are implemented within the building."},
            {"text": "My organization has measures in place to manage the operating costs associated with unreliable power supply (e.g., load management systems, efficient generator use)."},
            {"text": "Building operations in my organization are aligned with recognized green or sustainable building standards."},
        ],
    },
    {
        "code": "D",
        "title": "SECTION D: OFFICE BUILDING PERFORMANCE",
        "items": [
            {"text": "The building's spatial layout and infrastructure effectively support the core operations of my organization and can be reconfigured as work patterns change."},
            {"text": "The building's structural, electrical, and mechanical systems operate reliably, and support/maintenance services are delivered promptly and efficiently."},
            {"text": "My organization's occupancy costs are well managed, and the building's asset value is preserved in line with financial expectations."},
            {"text": "The building performs well in terms of energy efficiency, waste management, and overall environmental impact."},
            {"text": "Occupants experience a comfortable working environment (thermal, acoustic, and lighting conditions) and are generally satisfied with facility services."},
            {"text": "The building's physical assets are maintained in a manner that preserves their long-term value across the building's lifecycle."},
        ],
    },
    {
        "code": "E",
        "title": "SECTION E: LEVEL OF INTEGRATION OF FM PRACTICES WITHIN ORGANIZATIONAL STRATEGIC PLANNING AND ITS EFFECT ON BUILDING PERFORMANCE",
        "items": [
            {"text": "Facilities management is represented at board or senior management level in strategic decision-making."},
            {"text": "FM budgets and planning are aligned with my organization's overall strategic objectives."},
            {"text": "Facility managers are involved in strategic planning processes from an early stage, rather than only after major decisions have already been made."},
            {"text": "There are clear communication channels or formal mechanisms connecting FM decision-making to the organization's strategic planning process."},
            {"text": "My organization evaluates facility decisions (e.g., on space, maintenance, or capital upgrades) based on their contribution to organizational performance, not merely on cost."},
            {"text": "FM is treated as a value-adding function in my organization rather than simply a cost to be minimized."},
        ],
    },
    {
        "code": "F",
        "title": "SECTION F: MAJOR BARRIERS AFFECTING THE EFFECTIVE IMPLEMENTATION OF FM PRACTICES",
        "items": [
            {"text": "Inadequate funding or budgetary allocation limits the effective execution of FM practices in my organization."},
            {"text": "The high cost of technology, equipment, and skilled personnel constrains investment in advanced FM practices."},
            {"text": "Facilities management in my organization lacks sufficiently strong leadership commitment or support from senior management."},
            {"text": "Weak enforcement of facilities management regulations or standards hampers effective implementation."},
            {"text": "Limited adoption of FM technology (e.g., CMMS, CAFM, BIM) constrains effective facilities management."},
            {"text": "Facilities management personnel lack adequate training and professional capacity development opportunities."},
        ],
    },
]

SCALE = ["SA", "A", "N", "D", "SD"]
SCALE_LABELS = {
    "SA": "Strongly Agree",
    "A": "Agree",
    "N": "Neutral",
    "D": "Disagree",
    "SD": "Strongly Disagree",
}

# Build flat list of Likert field names: e.g. B1..B7, C1..C25 ...
def all_fields():
    fields = []
    for sec in SECTIONS:
        for i in range(1, len(sec["items"]) + 1):
            fields.append(f"{sec['code']}{i}")
    return fields

FIELDS = all_fields()
DEMO_FIELDS = [d["field"] for d in DEMOGRAPHICS]

# ---------------------------------------------------------------------------
# DB helpers
# ---------------------------------------------------------------------------
def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB)
        g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()

def init_db():
    conn = sqlite3.connect(DB)
    demo_cols = ", ".join(f'"{f}" TEXT' for f in DEMO_FIELDS)
    likert_cols = ", ".join(f'"{f}" TEXT' for f in FIELDS)
    conn.execute(f'''
        CREATE TABLE IF NOT EXISTS responses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            submitted_at TEXT,
            {demo_cols},
            {likert_cols}
        )
    ''')
    conn.commit()
    conn.close()

# ---------------------------------------------------------------------------
# Templates
# ---------------------------------------------------------------------------
BASE_CSS = """
<style>
  body { font-family: Arial, Helvetica, sans-serif; background:#f4f6f8; margin:0; padding:0; }
  .wrap { max-width: 800px; margin: 30px auto; background:#fff; border-radius:8px;
          box-shadow:0 1px 4px rgba(0,0,0,0.15); overflow:hidden; }
  .header { background:#2f6f4f; color:#fff; padding:24px 32px; }
  .header h1 { margin:0 0 6px 0; font-size:22px; }
  .header p { margin:0; font-size:14px; opacity:0.9; }
  .body { padding:24px 32px; }
  .intro { background:#fff8e1; border-left:4px solid #f0ad4e; padding:12px 16px; margin-bottom:20px; font-size:14px; }
  fieldset { border:1px solid #ddd; border-radius:6px; margin-bottom:20px; padding:16px; }
  legend { font-weight:bold; color:#2f6f4f; padding:0 8px; }
  .q { margin-bottom:14px; }
  .q label.qtext { display:block; font-weight:500; margin-bottom:6px; }
  .options { display:flex; gap:16px; flex-wrap:wrap; }
  .options label { font-weight:normal; display:flex; align-items:center; gap:4px; cursor:pointer; }
  select, input[type=text] { padding:6px 8px; border:1px solid #ccc; border-radius:4px; width:100%; max-width:300px; }
  .scale-table { width:100%; border-collapse: collapse; font-size:14px;}
  .scale-table th, .scale-table td { padding:8px; text-align:left; border-bottom:1px solid #eee; vertical-align:middle; }
  .scale-table td.item { text-align:left; }
  .scale-table td.pick { text-align:right; white-space:nowrap; }
  .scale-select { padding:6px 8px; border:1px solid #ccc; border-radius:4px; min-width:160px; }
  .scale-select:invalid { color:#888; }
  .submit-btn { background:#2f6f4f; color:#fff; border:none; padding:12px 28px; border-radius:6px;
                font-size:16px; cursor:pointer; }
  .submit-btn:hover { background:#255a3f; }
  .footer-links { text-align:center; margin: 16px 0 30px; font-size: 13px; }
  .footer-links a { color:#2f6f4f; }
  .success { text-align:center; padding: 60px 20px; }
  .success h2 { color:#2f6f4f; }
  .required { color:#c0392b; }
  .admin-bar { max-width:800px; margin: 10px auto 0; text-align:right; }
  .admin-bar a { font-size:13px; color:#2f6f4f; text-decoration:none; }
  table.data { border-collapse: collapse; width:100%; font-size:12px; }
  table.data th, table.data td { border:1px solid #ddd; padding:4px 6px; }
  table.data th { background:#2f6f4f; color:#fff; position: sticky; top:0; }
</style>
"""

FORM_TEMPLATE = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{{ title }}</title>
  """ + BASE_CSS + """
</head>
<body>
<div class="wrap">
  <div class="header">
    <h1>{{ title }}</h1>
    <p>{{ subtitle }}</p>
  </div>
  <div class="body">
    <div class="intro">
      Dear Respondent, {{ intro_text }}
    </div>

    <form method="POST" action="{{ url_for('submit') }}">
      <fieldset>
        <legend>SECTION A: DEMOGRAPHIC INFORMATION</legend>

        {% for demo in demographics %}
        <div class="q">
          <label class="qtext">{{ demo.label }} <span class="required">*</span></label>
          <select class="scale-select" name="{{ demo.field }}" required>
            <option value="" disabled selected>Select&hellip;</option>
            {% for opt in demo.options %}
            <option value="{{ opt }}">{{ opt }}</option>
            {% endfor %}
          </select>
        </div>
        {% endfor %}
      </fieldset>

      <div class="intro" style="background:#eef5ff;border-left-color:#3b82f6;">
        For Sections B to F, please select as appropriate:<br>
        SA = Strongly Agree, A = Agree, N = Neutral, D = Disagree, SD = Strongly Disagree.
      </div>

      {% for sec in sections %}
      <fieldset>
        <legend>{{ sec.title }}</legend>
        <table class="scale-table">
          {% for item in sec['items'] %}
          {% if item.group %}
          <tr><td colspan="2" style="padding-top:16px;"><strong style="color:#2f6f4f;">{{ item.group }}</strong></td></tr>
          {% endif %}
          <tr>
            <td class="item">{{ loop.index }}. {{ item.text }}</td>
            <td class="pick">
              <select class="scale-select" name="{{ sec.code }}{{ loop.index }}" required>
                <option value="" disabled selected>Select&hellip;</option>
                {% for s in scale %}
                <option value="{{ s }}">{{ s }} &mdash; {{ scale_labels[s] }}</option>
                {% endfor %}
              </select>
            </td>
          </tr>
          {% endfor %}
        </table>
      </fieldset>
      {% endfor %}

      <div style="text-align:center;">
        <button type="submit" class="submit-btn">Submit Questionnaire</button>
      </div>
    </form>
  </div>
</div>
</body>
</html>
"""

SUCCESS_TEMPLATE = """
<!doctype html>
<html><head><meta charset="utf-8">""" + BASE_CSS + """</head>
<body>
<div class="wrap"><div class="success">
  <h2>Thank you!</h2>
  <p>Your response has been recorded successfully.</p>
  <p><a href="{{ url_for('form') }}">Submit another response</a></p>
</div></div>
</body></html>
"""

ADMIN_TEMPLATE = """
<!doctype html>
<html><head><meta charset="utf-8">""" + BASE_CSS + """</head>
<body>
<div class="wrap" style="max-width:95%;">
  <div class="header"><h1>Admin — Responses</h1><p>{{ count }} response(s) recorded</p></div>
  <div class="body">
    <p>
      <a href="{{ url_for('export') }}" class="submit-btn" style="text-decoration:none;display:inline-block;">Export to Excel</a>
      &nbsp; <a href="{{ url_for('form') }}">Back to public form</a>
    </p>
    <div style="overflow:auto; max-height:70vh;">
    <table class="data">
      <tr>{% for h in headers %}<th>{{ h }}</th>{% endfor %}</tr>
      {% for row in rows %}
      <tr>{% for c in row %}<td>{{ c }}</td>{% endfor %}</tr>
      {% endfor %}
    </table>
    </div>
  </div>
</div>
</body></html>
"""

# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/input")
def form():
    return render_template_string(
        FORM_TEMPLATE,
        title=TITLE,
        subtitle=SUBTITLE,
        intro_text=INTRO_TEXT,
        demographics=DEMOGRAPHICS,
        sections=SECTIONS,
        scale=SCALE,
        scale_labels=SCALE_LABELS,
    )

@app.route("/submit", methods=["POST"])
def submit():
    db = get_db()
    demo_values = [request.form.get(f, "") for f in DEMO_FIELDS]
    values = [request.form.get(f, "") for f in FIELDS]

    placeholders = ", ".join(["?"] * (1 + len(DEMO_FIELDS) + len(FIELDS)))
    cols = ", ".join(
        ['"submitted_at"']
        + [f'"{f}"' for f in DEMO_FIELDS]
        + [f'"{f}"' for f in FIELDS]
    )
    db.execute(
        f'INSERT INTO responses ({cols}) VALUES ({placeholders})',
        [datetime.now().isoformat()] + demo_values + values,
    )
    db.commit()
    return redirect(url_for("success"))

@app.route("/success")
def success():
    return render_template_string(SUCCESS_TEMPLATE)

@app.route("/view")
def admin():
    db = get_db()
    cur = db.execute("SELECT * FROM responses ORDER BY id DESC")
    rows = cur.fetchall()
    headers = rows[0].keys() if rows else ["id", "submitted_at"] + DEMO_FIELDS + FIELDS
    return render_template_string(ADMIN_TEMPLATE, headers=headers, rows=rows, count=len(rows))

@app.route("/view/export")
def export():
    if openpyxl is None:
        return "openpyxl not installed. Run: pip install openpyxl", 500
    db = get_db()
    cur = db.execute("SELECT * FROM responses ORDER BY id")
    rows = cur.fetchall()
    headers = list(rows[0].keys()) if rows else ["id", "submitted_at"] + DEMO_FIELDS + FIELDS

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Responses"
    ws.append(headers)
    for row in rows:
        ws.append([row[h] for h in headers])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="questionnaire_responses.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

if __name__ == "__main__":
    init_db()
    app.run(debug=True, host="0.0.0.0", port=5000)
else:
    # Under a WSGI server (e.g. PythonAnywhere, gunicorn) this module is
    # imported rather than run directly, so make sure the table exists here too.
    init_db()